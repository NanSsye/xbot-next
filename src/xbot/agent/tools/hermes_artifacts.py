"""Safe user-facing artifact creation for the embedded Hermes runtime."""

from __future__ import annotations

import contextvars
import json
import os
import re
import sys
import uuid
from html import escape
from pathlib import Path
from typing import Any

_TOOLSET = "artifacts"
_MAX_CONTENT_CHARS = 120_000
_MAX_TABLE_CELLS = 4_000
_MAX_SLIDES = 50
_MAX_OUTPUT_BYTES = 30 * 1024 * 1024
_EXTENSIONS = {"md": ".md", "docx": ".docx", "xlsx": ".xlsx", "pptx": ".pptx", "pdf": ".pdf"}
_ARTIFACT_CONTEXT: contextvars.ContextVar[dict[str, Any] | None] = contextvars.ContextVar(
    "xbot_artifact_context", default=None
)


def set_artifact_context(context: dict[str, Any]):
    return _ARTIFACT_CONTEXT.set(context)


def reset_artifact_context(token) -> None:
    _ARTIFACT_CONTEXT.reset(token)


def _error(code: str, message: str) -> str:
    return json.dumps({"success": False, "error": code, "message": message}, ensure_ascii=False)


def _safe_filename(value: object, artifact_format: str) -> str:
    extension = _EXTENSIONS[artifact_format]
    raw = Path(str(value or f"document{extension}")).name
    stem = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", Path(raw).stem).strip(" ._")
    return f"{(stem or 'document')[:80]}{extension}"


def _content(args: dict[str, Any]) -> str:
    value = str(args.get("content") or "")
    if len(value) > _MAX_CONTENT_CHARS:
        raise ValueError(f"content exceeds {_MAX_CONTENT_CHARS} characters")
    return value


def _table(args: dict[str, Any]) -> list[list[Any]]:
    value = args.get("table") or []
    if not isinstance(value, list):
        raise ValueError("table must be an array of rows")
    rows: list[list[Any]] = []
    cells = 0
    for raw_row in value[:500]:
        if not isinstance(raw_row, list):
            raise ValueError("each table row must be an array")
        row = list(raw_row[:100])
        cells += len(row)
        if cells > _MAX_TABLE_CELLS:
            raise ValueError(f"table exceeds {_MAX_TABLE_CELLS} cells")
        rows.append(row)
    return rows


def _slides(args: dict[str, Any]) -> list[dict[str, Any]]:
    value = args.get("slides") or []
    if not isinstance(value, list):
        raise ValueError("slides must be an array")
    slides = [item for item in value if isinstance(item, dict)][:_MAX_SLIDES]
    for item in slides:
        bullets = item.get("bullets") or []
        if not isinstance(bullets, list) or len(bullets) > 30:
            raise ValueError("each slide may contain at most 30 bullets")
    return slides


def _lines(content: str) -> list[str]:
    return content.replace("\r\n", "\n").replace("\r", "\n").split("\n")


def _ensure_dependency_path() -> None:
    hermes_home = Path(os.environ.get("HERMES_HOME") or "data/hermes")
    dependency_path = str((hermes_home / "python-packages").resolve())
    if Path(dependency_path).is_dir() and dependency_path not in sys.path:
        sys.path.insert(0, dependency_path)


def _create_markdown(path: Path, *, title: str, content: str, table: list[list[Any]]) -> None:
    parts: list[str] = [f"# {title}"] if title else []
    if content:
        parts.append(content)
    if table:
        width = max(len(row) for row in table)
        normalized = [[str(cell) for cell in row] + [""] * (width - len(row)) for row in table]
        parts.extend([
            "| " + " | ".join(normalized[0]) + " |",
            "| " + " | ".join(["---"] * width) + " |",
            *("| " + " | ".join(row) + " |" for row in normalized[1:]),
        ])
    path.write_text("\n\n".join(part for part in parts if part).rstrip() + "\n", encoding="utf-8")


def _create_docx(path: Path, *, title: str, content: str, table: list[list[Any]]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    if title:
        heading = document.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in _lines(content):
        stripped = line.strip()
        if not stripped:
            document.add_paragraph()
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            document.add_paragraph(stripped)
    if table:
        width = max(len(row) for row in table)
        word_table = document.add_table(rows=len(table), cols=width)
        word_table.style = "Table Grid"
        for row_index, row in enumerate(table):
            for column_index, cell in enumerate(row):
                word_table.cell(row_index, column_index).text = str(cell)
    document.save(path)


def _create_xlsx(path: Path, *, title: str, content: str, table: list[list[Any]], sheet_name: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_name or "Sheet1")[:31]
    row_index = 1
    if title:
        sheet.cell(row=row_index, column=1, value=title).font = Font(name="Arial", size=16, bold=True)
        row_index += 2
    rows = table or [[line] for line in _lines(content) if line.strip()]
    for relative_row, row in enumerate(rows, start=row_index):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=relative_row, column=column_index, value=value)
            cell.font = Font(name="Arial", size=10, bold=relative_row == row_index and bool(table))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if relative_row == row_index and table:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for column_index in range(1, max((len(row) for row in rows), default=1) + 1):
        values = [str(sheet.cell(row=row, column=column_index).value or "") for row in range(1, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max(map(len, values)), 10) + 2, 48)
    sheet.freeze_panes = f"A{row_index + 1}" if table else None
    workbook.save(path)


def _create_pptx(path: Path, *, title: str, content: str, slides: list[dict[str, Any]]) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    if title:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = str(slides[0].get("subtitle") or "") if slides else ""
    if not slides:
        chunks = [chunk.strip() for chunk in content.split("\n\n") if chunk.strip()]
        slides = [{"title": f"第 {index + 1} 页", "bullets": chunk.splitlines()} for index, chunk in enumerate(chunks)]
    for item in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
        slide.shapes.title.text = str(item.get("title") or "内容")[:200]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, bullet in enumerate(item.get("bullets") or []):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(bullet)[:1200]
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(22)
    presentation.save(path)


def _create_pdf(path: Path, *, title: str, content: str, table: list[list[Any]]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    normal = ParagraphStyle("ChineseNormal", fontName="STSong-Light", fontSize=10.5, leading=17)
    heading = ParagraphStyle("ChineseTitle", parent=normal, fontSize=18, leading=26, spaceAfter=10)
    story: list[Any] = []
    if title:
        story.extend([Paragraph(escape(title), heading), Spacer(1, 4 * mm)])
    for line in _lines(content):
        story.extend([Paragraph(escape(line) if line else "&nbsp;", normal), Spacer(1, 2 * mm)])
    if table:
        pdf_table = Table([[Paragraph(escape(str(cell)), normal) for cell in row] for row in table], repeatRows=1)
        pdf_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AAB7C4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(pdf_table)
    SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
    ).build(story)


def _send_created_file(path: Path, channel: str) -> dict[str, Any]:
    if channel == "qq":
        from xbot.agent.tools.hermes_qq import send_file
    elif channel == "telegram":
        from xbot.agent.tools.hermes_telegram import send_file
    else:
        from xbot.agent.tools.hermes_wechat import send_file
    try:
        raw_result = send_file({"path": str(path)})
        result = json.loads(raw_result) if isinstance(raw_result, str) else raw_result
    except Exception as exc:
        return {"sent": False, "send_error": str(exc)[:500]}
    if isinstance(result, dict) and not result.get("error"):
        return {"sent": bool(result.get("success", True))}
    return {"sent": False, "send_error": str(result)[:500]}


def create_artifact(args: dict[str, Any], **_: Any) -> str:
    context = _ARTIFACT_CONTEXT.get()
    if not context:
        return _error("artifact_context_unavailable", "当前会话的文件输出目录不可用。")
    artifact_format = str(args.get("format") or "").strip().lower() if isinstance(args, dict) else ""
    if artifact_format not in _EXTENSIONS:
        return _error("invalid_format", "仅支持 md、docx、xlsx、pptx、pdf。")
    try:
        _ensure_dependency_path()
        content = _content(args)
        table = _table(args)
        slides = _slides(args)
        title = str(args.get("title") or "").strip()[:300]
        filename = _safe_filename(args.get("filename"), artifact_format)
        output_dir = Path(str(context["output_dir"])).resolve()
        output_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
        path = output_dir / filename
        temporary = output_dir / f".{uuid.uuid4().hex}{_EXTENSIONS[artifact_format]}"
        creators = {
            "md": lambda: _create_markdown(temporary, title=title, content=content, table=table),
            "docx": lambda: _create_docx(temporary, title=title, content=content, table=table),
            "xlsx": lambda: _create_xlsx(temporary, title=title, content=content, table=table, sheet_name=str(args.get("sheet_name") or "Sheet1")),
            "pptx": lambda: _create_pptx(temporary, title=title, content=content, slides=slides),
            "pdf": lambda: _create_pdf(temporary, title=title, content=content, table=table),
        }
        creators[artifact_format]()
        size = temporary.stat().st_size
        if size <= 0 or size > _MAX_OUTPUT_BYTES:
            temporary.unlink(missing_ok=True)
            return _error("invalid_output", "生成文件为空或超过 30MB。")
        os.chmod(temporary, 0o600)
        os.replace(temporary, path)
        delivery = {"sent": False}
        if args.get("send", True):
            delivery = _send_created_file(path, str(context.get("channel") or "wechat"))
        return json.dumps({
            "success": True, "format": artifact_format, "filename": filename,
            "path": str(path), "size": size, **delivery,
        }, ensure_ascii=False)
    except (ImportError, ModuleNotFoundError) as exc:
        return _error("artifact_dependency_missing", f"{artifact_format} 生成组件未安装：{exc}")
    except Exception as exc:
        return _error("artifact_create_failed", str(exc)[:800])


def register_xbot_artifact_tools() -> None:
    from tools.registry import registry

    schema = {
        "name": "artifact_create",
        "description": "Create and optionally send a Markdown, Word, Excel, PowerPoint, or PDF file in the current user's isolated output directory.",
        "parameters": {
            "type": "object",
            "properties": {
                "format": {"type": "string", "enum": ["md", "docx", "xlsx", "pptx", "pdf"]},
                "filename": {"type": "string", "description": "User-facing filename; extension is normalized automatically."},
                "title": {"type": "string"},
                "content": {"type": "string", "description": "Main textual content."},
                "table": {"type": "array", "items": {"type": "array", "items": {}}},
                "slides": {"type": "array", "items": {"type": "object", "properties": {"title": {"type": "string"}, "subtitle": {"type": "string"}, "bullets": {"type": "array", "items": {"type": "string"}}}}},
                "sheet_name": {"type": "string"},
                "send": {"type": "boolean", "default": True},
            },
            "required": ["format", "filename"],
            "additionalProperties": False,
        },
    }
    existing = registry.get_entry("artifact_create")
    if existing is not None and existing.handler is create_artifact and existing.toolset == _TOOLSET:
        return
    registry.register(
        name="artifact_create", toolset=_TOOLSET, schema=schema, handler=create_artifact,
        description=schema["description"], override=existing is not None,
    )
