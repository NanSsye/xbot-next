from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_SUFFIXES = {
    ".txt", ".md", ".markdown", ".pdf", ".docx", ".xlsx", ".pptx",
    ".png", ".jpg", ".jpeg", ".webp", ".gif",
}
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_EXTRACT_CHARS = 80_000


def extract_attachment(source: Path, copied_to: Path) -> tuple[str, str]:
    source = source.resolve()
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        return "unsupported", ""
    if not source.is_file() or source.is_symlink():
        return "missing", ""
    if source.stat().st_size > MAX_FILE_BYTES:
        return "too_large", ""
    copied_to.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, copied_to)
    try:
        if suffix in {".txt", ".md", ".markdown"}:
            text = source.read_text(encoding="utf-8", errors="replace")
        elif suffix == ".pdf":
            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(source).pages)
        elif suffix == ".docx":
            text = "\n".join(p.text for p in Document(source).paragraphs if p.text.strip())
        elif suffix == ".xlsx":
            workbook = load_workbook(source, read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in workbook.worksheets[:20]:
                parts.append(f"## 工作表：{sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    line = "\t".join("" if value is None else str(value) for value in row)
                    if line.strip():
                        parts.append(line)
                    if sum(len(item) for item in parts) >= MAX_EXTRACT_CHARS:
                        break
            text = "\n".join(parts)
        elif suffix == ".pptx":
            deck = Presentation(source)
            parts = []
            for number, slide in enumerate(deck.slides, 1):
                lines = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
                if lines:
                    parts.append(f"## 第 {number} 页\n" + "\n".join(lines))
            text = "\n\n".join(parts)
        else:
            text = f"图片附件：{source.name}。原图已保存在 Vault raw/files 中，当前版本不执行图片内文字 OCR。"
    except Exception as exc:
        return "failed", f"提取失败：{type(exc).__name__}"
    return "ready", text[:MAX_EXTRACT_CHARS]
